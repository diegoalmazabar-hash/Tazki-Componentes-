# -*- coding: utf-8 -*-
"""Mapa de riesgos DS 44 - Tazki.
Construido sobre la "Guia para la elaboracion de Mapas de Riesgos Laborales al
interior de las entidades empleadoras", dictada por Resolucion Exenta N°129 de
2025 de la Subsecretaria de Prevision Social (art. 62 inciso final del DS 44),
vigente desde el 1 de marzo de 2025. Simbologia del Anexo 1; escalas del 6.2.1.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

NAVY, COBALT, LINE = "0C1E50", "1537D1", "C7D0E0"
SOFT, BAND, GREY, CALC = "EEF2FE", "F4F6FB", "6B76A0", "E8EEFB"
ROJO, NARANJA, AMARILLO, VERDE = "F4C7C3", "FBE0C8", "FFF2C2", "D7F0DF"
WARN = "FBE9E7"
FN = "Arial"
def fo(sz=10, b=False, c="0C1E50", it=False): return Font(name=FN, size=sz, bold=b, color=c, italic=it)
thin = Side(style="thin", color=LINE); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()
wb.remove(wb.active)   # quitar la hoja vacia por defecto
def banda(ws, rng, texto, alto=30):
    ws.merge_cells(rng); c = ws[rng.split(":")[0]]
    c.value=texto; c.font=fo(14,True,"FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY)
    c.alignment=Alignment(vertical="center",horizontal="left",indent=1)
    ws.row_dimensions[int("".join(x for x in rng.split(":")[0] if x.isdigit()))].height=alto
def sub(ws, rng, texto, alto=26):
    ws.merge_cells(rng); c = ws[rng.split(":")[0]]
    c.value=texto; c.font=fo(9,False,"FFFFFF"); c.fill=PatternFill("solid",fgColor=COBALT)
    c.alignment=Alignment(vertical="center",horizontal="left",indent=1,wrap_text=True)
    ws.row_dimensions[int("".join(x for x in rng.split(":")[0] if x.isdigit()))].height=alto
def heads(ws, headers, row, col0=2, alto=46):
    for i,h in enumerate(headers):
        c=ws.cell(row=row,column=col0+i,value=h); c.font=fo(8.5,True,"FFFFFF")
        c.fill=PatternFill("solid",fgColor=COBALT); c.border=box
        c.alignment=Alignment(vertical="center",horizontal="center",wrap_text=True)
    ws.row_dimensions[row].height=alto

SEG="Seguridad, emergencia y desastres"; HIG="Higiénico químico o físico"
MSK="Músculo-esquelético"; PSI="Psicosocial"
# (riesgo especifico, familia, escala, exige leyenda, que detallar, pictograma)
S = [
("Caídas al mismo nivel","Caída de personas",SEG,"No","","Persona tropezando o resbalando sobre una superficie"),
("Caídas a distinto nivel","Caída de personas",SEG,"No","","Persona cayendo por un desnivel o escalón"),
("Caídas de altura","Caída de personas",SEG,"No","","Persona cayendo desde un borde elevado"),
("Caídas al agua","Caída de personas",SEG,"No","","Persona cayendo a la superficie del agua"),
("Atrapamiento","Contacto con objetos",SEG,"No","","Mano entre engranaje y correa o rodillo"),
("Caída de objetos","Contacto con objetos",SEG,"No","","Carga suspendida cayendo sobre una persona"),
("Cortes por objetos o herramientas corto-punzantes","Contacto con objetos",SEG,"No","","Mano frente a hoja o sierra de corte"),
("Choque contra objetos","Contacto con objetos",SEG,"No","","Persona chocando contra un obstáculo a la altura de la cabeza"),
("Contacto con personas (agresión)","Contacto con seres vivos",SEG,"No","","Dos puños enfrentados"),
("Contacto con animales o insectos","Contacto con seres vivos",SEG,"No","","Perro en actitud de ataque"),
("Contactos térmicos por calor","Contactos térmicos",SEG,"Sí","Indicar si el contacto es por calor o por frío","Superficie caliente con líneas de calor · un solo símbolo para la familia"),
("Contactos térmicos por frío","Contactos térmicos",SEG,"Sí","Indicar si el contacto es por calor o por frío","Mismo símbolo de la familia"),
("Contactos eléctricos directos en baja tensión","Contacto con energía eléctrica",SEG,"Sí","Indicar si es directo o indirecto y si es baja o alta tensión","Rayo eléctrico · un solo símbolo para la familia"),
("Contactos eléctricos directos en alta tensión","Contacto con energía eléctrica",SEG,"Sí","Indicar si es directo o indirecto y si es baja o alta tensión","Mismo símbolo de la familia"),
("Contactos eléctricos indirectos en baja tensión","Contacto con energía eléctrica",SEG,"Sí","Indicar si es directo o indirecto y si es baja o alta tensión","Mismo símbolo de la familia"),
("Contactos eléctricos indirectos en alta tensión","Contacto con energía eléctrica",SEG,"Sí","Indicar si es directo o indirecto y si es baja o alta tensión","Mismo símbolo de la familia"),
("Contacto con sustancias cáusticas o corrosivas","Contacto con sustancias químicas",SEG,"Sí","Indicar la sustancia","Probeta derramando líquido sobre una mano · un solo símbolo para la familia"),
("Contacto con otras sustancias químicas","Contacto con sustancias químicas",SEG,"Sí","Indicar la sustancia","Mismo símbolo de la familia"),
("Explosiones","Contacto con elementos que se proyectan",SEG,"Sí","Indicar si es explosión o proyección de fragmentos","Estallido con fragmentos · un solo símbolo para la familia"),
("Proyección de fragmentos o partículas","Contacto con elementos que se proyectan",SEG,"Sí","Indicar si es explosión o proyección de fragmentos","Mismo símbolo de la familia"),
("Atropellos o golpes con vehículos","Vehículos en movimiento",SEG,"Sí","Indicar el tipo de evento","Vehículo industrial con flecha de movimiento"),
("Choque, colisión o volcamiento","Vehículos en movimiento",SEG,"Sí","Indicar el tipo de evento","Vehículo industrial con flecha de movimiento"),
("Incendios","Incendios",SEG,"No","","Llama"),
("Exposición a ambientes con deficiencia de oxígeno","Condiciones atmosféricas extremas",SEG,"No","","Persona desvanecida en un recinto cerrado"),
("Exposición a sustancias químicas tóxicas","Condiciones atmosféricas extremas",SEG,"Sí","Indicar la sustancia","Persona con máscara en atmósfera con partículas"),
("Ingesta de sustancias nocivas","Ingesta de sustancias nocivas",SEG,"No","","Persona con la mano en la boca y un recipiente"),
("Otros riesgos de seguridad","Otros riesgos",SEG,"Sí","Indicar de qué riesgo se trata","Signo de exclamación · peligro general"),
("Exposición a aerosoles sólidos","Exposición a agentes químicos",HIG,"Sí","Indicar el agente higiénico específico: sílice, polvo de madera, etc.","Persona con partículas alrededor de la cabeza"),
("Exposición a aerosoles líquidos","Exposición a agentes químicos",HIG,"Sí","Indicar el agente higiénico específico","Persona con gotas o niebla en las vías respiratorias"),
("Exposición a gases y vapores","Exposición a agentes químicos",HIG,"Sí","Indicar el agente higiénico específico","Persona con nube de gas o vapor"),
("Exposición a ruido","Exposición a agentes físicos",HIG,"No","","Oreja con ondas sonoras"),
("Exposición a vibraciones","Exposición a agentes físicos",HIG,"Sí","Indicar si la exposición es de cuerpo entero o mano-brazo","Mano o brazo con líneas de vibración"),
("Exposición a radiaciones ionizantes","Exposición a agentes físicos",HIG,"No","","Trébol radiactivo · ÚNICO símbolo que puede ir en MAGENTA, según la OIEA"),
("Exposición a radiaciones no ionizantes","Exposición a agentes físicos",HIG,"No","","Antena emitiendo ondas"),
("Exposición a calor","Exposición a agentes físicos",HIG,"No","","Termómetro con temperatura alta"),
("Exposición a frío","Exposición a agentes físicos",HIG,"No","","Copo de nieve"),
("Exposición a altas presiones","Exposición a agentes físicos",HIG,"No","","Escafandra · trabajo hiperbárico"),
("Exposición a bajas presiones","Exposición a agentes físicos",HIG,"No","","Persona en montañas · altura geográfica"),
("Transmisión por fluidos corporales","Exposición a peligros biológicos",SEG,"Sí","Indicar la vía de transmisión","Símbolo internacional de riesgo biológico · un solo símbolo para la familia"),
("Transmisión por inhalación, dermal, oral o parenteral","Exposición a peligros biológicos",SEG,"Sí","Indicar la vía de transmisión","Mismo símbolo de la familia"),
("Manejo o manipulación manual de cargas (MMC)","Músculo-esquelético",MSK,"Sí","Indicar la familia de riesgo músculo-esquelético","Persona agachada levantando una carga · un solo símbolo para toda la categoría"),
("Manejo o manipulación manual de personas o pacientes (MMP)","Músculo-esquelético",MSK,"Sí","Indicar la familia de riesgo músculo-esquelético","Mismo símbolo de la categoría"),
("Trabajo repetitivo de miembros superiores","Músculo-esquelético",MSK,"Sí","Indicar la familia de riesgo músculo-esquelético","Mismo símbolo de la categoría"),
("Posturas forzadas","Músculo-esquelético",MSK,"Sí","Indicar la familia de riesgo músculo-esquelético","Mismo símbolo de la categoría"),
("Posturas estáticas","Músculo-esquelético",MSK,"Sí","Indicar la familia de riesgo músculo-esquelético","Mismo símbolo de la categoría"),
("Riesgo psicosocial","Riesgos psicosociales",PSI,"No","","Cabeza con rayos alrededor · SIEMPRE va en el mapa, sea cual sea su nivel"),
]
NIV = ["Nivel 4","Nivel 3","Nivel 2","Nivel 1","Intolerable","Importante","Moderado","Tolerable",
       "Alto o crítico","Medio o razonable","Bajo"]

# ══════ HOJA 2 (se crea antes para poder referenciarla) · SIMBOLOGÍA ══════
wsS = wb.create_sheet("Simbología")
wsS.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",50),("C",34),("D",30),("E",13),("F",48),("G",54)]: wsS.column_dimensions[k].width=v
banda(wsS,"B2:G2","SIMBOLOGÍA · ANEXO 1 DE LA GUÍA")
sub(wsS,"B3:G3","Todas las señales son triángulos amarillos de precaución, con borde y pictograma negros. "
                "La única excepción es radiaciones ionizantes, que puede ir en magenta. Esta hoja alimenta las "
                "listas de la hoja «Inventario del mapa»: no borres ni reordenes sus filas.")
heads(wsS,["Riesgo específico","Familia de riesgos","Escala que le corresponde",
           "¿Detallar en la leyenda?","Qué hay que detallar","Pictograma"],5,alto=40)
sr0=6
for i,fila in enumerate(S):
    r=sr0+i
    for j,v in enumerate(fila):
        c=wsS.cell(row=r,column=2+j,value=v); c.border=box; c.font=fo(9)
        c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="center" if j==3 else "left")
        if j==3 and v=="Sí": c.fill=PatternFill("solid",fgColor=SOFT); c.font=fo(9,True)
    wsS.row_dimensions[r].height=26
srN=sr0+len(S)-1
wsS.freeze_panes="C6"; wsS.auto_filter.ref="B5:G%d"%srN

# ══════ HOJA 3 · NIVELES Y COLORES ══════
wsN = wb.create_sheet("Niveles y colores")
wsN.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",34),("C",26),("D",18),("E",64)]: wsN.column_dimensions[k].width=v
banda(wsN,"B2:E2","NIVELES DE RIESGO Y SU COLOR · PUNTO 6.2.1 DE LA GUÍA")
sub(wsN,"B3:E3","Ojo: la guía usa CUATRO escalas distintas según el tipo de riesgo, y no todas tienen los mismos "
                "niveles ni los mismos colores. Es el punto que más se equivoca al hacer un mapa.")
heads(wsN,["Tipo de riesgo","Nivel","Color del círculo","Nota"],5,alto=30)
ESC=[(HIG,"Nivel 4","Rojo","Riesgos higiénicos de origen químico y físico. Cuatro niveles numerados."),
     (HIG,"Nivel 3","Anaranjado",""),(HIG,"Nivel 2","Amarillo",""),(HIG,"Nivel 1","Verde",""),
     (SEG,"Intolerable","Rojo","Riesgos de seguridad, de emergencia y desastres, y peligros biológicos."),
     (SEG,"Importante","Anaranjado",""),(SEG,"Moderado","Amarillo",""),(SEG,"Tolerable","Verde",""),
     (MSK,"Alto o crítico","Rojo","Solo TRES niveles. El nivel medio va en AMARILLO, no anaranjado."),
     (MSK,"Medio o razonable","Amarillo",""),(MSK,"Bajo","Verde",""),
     (PSI,"Alto o crítico","Rojo","Solo TRES niveles. Acá el nivel medio va en ANARANJADO, no amarillo. Y el riesgo psicosocial SIEMPRE se incluye en el mapa, sea cual sea su nivel."),
     (PSI,"Medio o razonable","Anaranjado",""),(PSI,"Bajo","Verde","")]
COLF={"Rojo":ROJO,"Anaranjado":NARANJA,"Amarillo":AMARILLO,"Verde":VERDE}
nr0=6
for i,(t,n,col,nota) in enumerate(ESC):
    r=nr0+i
    for j,v in enumerate([t,n,col,nota]):
        c=wsN.cell(row=r,column=2+j,value=v); c.border=box; c.font=fo(9)
        c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="center" if j in (1,2) else "left")
        if j==2: c.fill=PatternFill("solid",fgColor=COLF[col]); c.font=fo(9,True)
    wsN.row_dimensions[r].height=26
nrN=nr0+len(ESC)-1
wsN.cell(row=nrN+2,column=2,value="Fuente: punto 6.2.1 de la Guía dictada por la Resolución Exenta N°129 de 2025.").font=fo(8.5,False,GREY)
# lista de niveles para la validacion, en una columna auxiliar
for i,n in enumerate(NIV):
    wsN.cell(row=nr0+i,column=8,value=n).font=fo(8,False,"FFFFFF")
wsN.column_dimensions["H"].hidden=True

# ══════ HOJA 1 · INVENTARIO DEL MAPA ══════
ws = wb.create_sheet("Inventario del mapa", 0)
ws.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",6),("C",24),("D",46),("E",28),("F",26),("G",18),("H",14),("I",30),
            ("J",13),("K",44),("L",34),("M",13),("N",30)]:
    ws.column_dimensions[k].width=v
banda(ws,"B2:N2","INVENTARIO DE RIESGOS DEL MAPA · DS 44")
sub(ws,"B3:N3","Este es el paso previo al dibujo: decidir qué riesgo va en qué zona, con qué color y qué hay que "
               "aclarar en la leyenda. Elige el riesgo y el nivel; la familia, la escala, el color y el aviso de "
               "leyenda se completan solos.")
H=["N°","Zona del croquis","Riesgo específico","Familia","Escala que corresponde","Nivel evaluado",
   "Color del círculo","Control","¿Detallar en la leyenda?","Qué hay que detallar",
   "Texto para la leyenda","¿Transversal?","Observaciones"]
heads(ws,H,5)
r0=6; N=120
SIM="Simbología"
for i in range(N):
    r=r0+i
    ws.cell(row=r,column=2,value=i+1)
    ws.cell(row=r,column=5,value='=IFERROR(INDEX(%s!$C$%d:$C$%d,MATCH($D%d,%s!$B$%d:$B$%d,0)),"")'%(SIM,sr0,srN,r,SIM,sr0,srN))
    ws.cell(row=r,column=6,value='=IFERROR(INDEX(%s!$D$%d:$D$%d,MATCH($D%d,%s!$B$%d:$B$%d,0)),"")'%(SIM,sr0,srN,r,SIM,sr0,srN))
    ws.cell(row=r,column=8,value=('=IF($G%d="","",IF(OR($G%d="Nivel 4",$G%d="Intolerable",$G%d="Alto o crítico"),"Rojo",'
        'IF(OR($G%d="Nivel 3",$G%d="Importante"),"Anaranjado",'
        'IF($G%d="Medio o razonable",IF($F%d="Psicosocial","Anaranjado","Amarillo"),'
        'IF(OR($G%d="Nivel 2",$G%d="Moderado"),"Amarillo",'
        'IF(OR($G%d="Nivel 1",$G%d="Tolerable",$G%d="Bajo"),"Verde",""))))))')%((r,)*13))
    ws.cell(row=r,column=9,value=('=IF(OR($F%d="",$G%d=""),"",IF(OR('
        'AND($F%d="Higiénico químico o físico",OR($G%d="Nivel 4",$G%d="Nivel 3",$G%d="Nivel 2",$G%d="Nivel 1")),'
        'AND($F%d="Seguridad, emergencia y desastres",OR($G%d="Intolerable",$G%d="Importante",$G%d="Moderado",$G%d="Tolerable")),'
        'AND(OR($F%d="Músculo-esquelético",$F%d="Psicosocial"),OR($G%d="Alto o crítico",$G%d="Medio o razonable",$G%d="Bajo"))'
        '),"OK","⚠ El nivel no corresponde a esta escala"))')%((r,)*17))
    ws.cell(row=r,column=10,value='=IFERROR(INDEX(%s!$E$%d:$E$%d,MATCH($D%d,%s!$B$%d:$B$%d,0)),"")'%(SIM,sr0,srN,r,SIM,sr0,srN))
    ws.cell(row=r,column=11,value='=IFERROR(INDEX(%s!$F$%d:$F$%d,MATCH($D%d,%s!$B$%d:$B$%d,0)),"")'%(SIM,sr0,srN,r,SIM,sr0,srN))
    for col in range(2,15):
        c=ws.cell(row=r,column=col); c.border=box; c.font=fo(9)
        c.alignment=Alignment(vertical="top",wrap_text=True,
                              horizontal="center" if col in (2,8,9,10,14-0) else "left")
        if col in (5,6,8,9,10,11): c.fill=PatternFill("solid",fgColor=CALC)
        elif col in (3,4,7,12,13,14): c.fill=PatternFill("solid",fgColor=BAND)
    ws.cell(row=r,column=13).alignment=Alignment(horizontal="center",vertical="top")
    ws.row_dimensions[r].height=30
rN=r0+N-1
DataValidation
dv1=DataValidation(type="list",allow_blank=True,formula1="=%s!$B$%d:$B$%d"%(SIM,sr0,srN)); ws.add_data_validation(dv1); dv1.add("D%d:D%d"%(r0,rN))
dv2=DataValidation(type="list",allow_blank=True,formula1="='Niveles y colores'!$H$%d:$H$%d"%(nr0,nr0+len(NIV)-1)); ws.add_data_validation(dv2); dv2.add("G%d:G%d"%(r0,rN))
dv3=DataValidation(type="list",allow_blank=True,formula1='"Sí,No"'); ws.add_data_validation(dv3); dv3.add("M%d:M%d"%(r0,rN))
for nombre,color in [("Rojo",ROJO),("Anaranjado",NARANJA),("Amarillo",AMARILLO),("Verde",VERDE)]:
    ws.conditional_formatting.add("H%d:H%d"%(r0,rN),
        FormulaRule(formula=['$H%d="%s"'%(r0,nombre)],fill=PatternFill("solid",fgColor=color)))
ws.conditional_formatting.add("I%d:I%d"%(r0,rN),
    FormulaRule(formula=['LEFT($I%d,1)="⚠"'%r0],fill=PatternFill("solid",fgColor=WARN),font=fo(9,True,"C0392B")))
ws.conditional_formatting.add("J%d:J%d"%(r0,rN),
    FormulaRule(formula=['$J%d="Sí"'%r0],fill=PatternFill("solid",fgColor=SOFT),font=fo(9,True)))
ws.freeze_panes="D6"; ws.auto_filter.ref="B5:N%d"%rN

# ══════ HOJA 4 · REGISTRO DE MAPAS ══════
wsR = wb.create_sheet("Registro de mapas")
wsR.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",6),("C",30),("D",38),("E",16),("F",22),("G",34),("H",16),("I",30)]:
    wsR.column_dimensions[k].width=v
banda(wsR,"B2:I2","REGISTRO DE MAPAS · UBICACIÓN, VIGENCIA Y PARTICIPACIÓN")
sub(wsR,"B3:I3","Un mapa por lugar de trabajo. La guía obliga a registrar la ubicación física de cada mapa en el "
                "Programa de Trabajo Preventivo, y a actualizarlo cada vez que cambie la matriz.")
heads(wsR,["N°","Lugar de trabajo","Ubicación física del mapa","Fecha de elaboración",
           "Versión de la matriz que refleja","Quién participó en la confección",
           "Última actualización","Observaciones"],5,alto=40)
for i in range(30):
    r=6+i
    wsR.cell(row=r,column=2,value=i+1)
    for col in range(2,10):
        c=wsR.cell(row=r,column=col); c.border=box; c.font=fo(9)
        c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="center" if col in (2,5,8) else "left")
        if col>2: c.fill=PatternFill("solid",fgColor=BAND)
    wsR.cell(row=r,column=5).number_format="dd-mm-yyyy"
    wsR.cell(row=r,column=8).number_format="dd-mm-yyyy"
    wsR.row_dimensions[r].height=28
dvR=DataValidation(type="list",allow_blank=True,
    formula1='"Comité Paritario,Delegado de SST,Departamento de Prevención,Organismo administrador,Comité Paritario y OAL,Delegado de SST y OAL"')
wsR.add_data_validation(dvR); dvR.add("G6:G35")

# ══════ HOJA 5 · CÓMO USARLA ══════
wsC = wb.create_sheet("Cómo usarla")
wsC.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",30),("C",104)]: wsC.column_dimensions[k].width=v
banda(wsC,"B2:C2","CÓMO USAR ESTA PLANTILLA")
F=[("Qué es y qué no es","Esta planilla NO dibuja el mapa. Hace el paso previo, que es donde se equivoca casi todo el mundo: decidir qué riesgo va en qué zona, de qué color va su círculo y qué hay que aclarar en la leyenda. El dibujo lo haces después, en el croquis, con estas decisiones ya tomadas."),
   ("De dónde sale","De la «Guía para la elaboración de Mapas de Riesgos Laborales al interior de las entidades empleadoras», dictada por la Resolución Exenta N°129 de 2025 de la Subsecretaría de Previsión Social, que el artículo 62 del DS 44 fija como metodología. Rige desde el 1 de marzo de 2025."),
   ("Paso 1 · El croquis","Elabora un plano o croquis del lugar de trabajo con sus zonas, y las máquinas y equipos si hacen falta. No tiene que ser exacto: tiene que ser claro y reconocible. Puedes reutilizar un croquis que ya exista."),
   ("Paso 2 · Los símbolos","Cada riesgo se marca con un CÍRCULO del color de su nivel de riesgo, y dentro va el TRIÁNGULO amarillo de precaución con el pictograma. Es esa combinación, no el triángulo solo. Los riesgos y niveles salen de tu matriz."),
   ("El detalle que casi nadie acierta","La guía usa cuatro escalas distintas según el tipo de riesgo, y no coinciden. En músculo-esquelético el nivel medio va en AMARILLO; en psicosocial el mismo nivel medio va en ANARANJADO. La hoja «Niveles y colores» las tiene las cuatro, y la columna «Control» del inventario te avisa si elegiste un nivel que no corresponde a esa escala."),
   ("Qué se puede priorizar","Si en un mismo lugar hay muchos riesgos, se pueden priorizar los de nivel rojo y anaranjado por sobre los amarillos y verdes. La excepción es el riesgo psicosocial: va siempre, sea cual sea su nivel."),
   ("Dónde va cada cosa en el croquis","Un riesgo transversal a todo el lugar va en la parte superior IZQUIERDA. La leyenda va en el sector inferior DERECHO. Ahí es donde se aclara el riesgo específico cuando la familia tiene un solo símbolo, el agente higiénico concreto, el tipo de vibración y la familia músculo-esquelética."),
   ("Dónde se cuelga y cada cuánto se revisa","En sitios visibles de cada lugar de trabajo, accesible para las personas trabajadoras. Se actualiza cada vez que cambie la matriz. Y la ubicación física de cada mapa debe quedar registrada en el Programa de Trabajo Preventivo: para eso está la hoja «Registro de mapas»."),
   ("Quién participa","El Comité Paritario participa en la confección. Si no hay Comité, participa el Delegado de Seguridad y Salud en el Trabajo. Se puede pedir asistencia técnica al Departamento de Prevención o al organismo administrador."),
   ("Personas especialmente sensibles","La guía recomienda ajustes para quienes no puedan leer el mapa como está: señalización en braille, dispositivos de audio o textos descriptivos que interpreten los colores. Aplica a personas con discapacidad visual, con problemas para discriminar colores y a personas neurodivergentes."),
   ("Una advertencia sobre la propia guía","El ejemplo del Anexo 2 de la guía oficial usa una leyenda de colores que NO coincide con la escala de su propio punto 6.2.1: en el ejemplo aparece rojo/amarillo/celeste y mezcla las nomenclaturas. Es una inconsistencia del documento. Esta planilla sigue el punto 6.2.1, que es el cuerpo normativo de la guía, y no el ejemplo."),
   ("Quién lo hizo","Tazki, software chileno de gestión de seguridad y salud en el trabajo. Si quieres que la matriz, el mapa y sus actualizaciones dejen de vivir en carpetas sueltas, escríbenos en tazki.cl")]
r=4
for t,d in F:
    a=wsC.cell(row=r,column=2,value=t); a.font=fo(10,True); a.fill=PatternFill("solid",fgColor=SOFT)
    a.alignment=Alignment(vertical="top",wrap_text=True); a.border=box
    b=wsC.cell(row=r,column=3,value=d); b.font=fo(10); b.alignment=Alignment(vertical="top",wrap_text=True); b.border=box
    wsC.row_dimensions[r].height=62; r+=1

# ══════ HOJA 6 · CONTROL DE VERSIONES ══════
wsV = wb.create_sheet("Control de versiones")
wsV.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",14),("C",16),("D",30),("E",56)]: wsV.column_dimensions[k].width=v
banda(wsV,"B2:E2","CONTROL DE VERSIONES")
heads(wsV,["Versión","Fecha","Responsable","Qué cambió"],4,alto=26)
for i in range(15):
    r=5+i
    for col in range(2,6):
        c=wsV.cell(row=r,column=col); c.border=box; c.font=fo(9)
        c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="center" if col in (2,3) else "left")
    wsV.cell(row=r,column=3).number_format="dd-mm-yyyy"; wsV.row_dimensions[r].height=24

for s in wb.worksheets: s.sheet_properties.tabColor=COBALT
wb.save("/tmp/piezas/Mapa-de-Riesgos-DS44-Tazki.xlsx")
print("guardado · hojas:", [s.title for s in wb.worksheets])
