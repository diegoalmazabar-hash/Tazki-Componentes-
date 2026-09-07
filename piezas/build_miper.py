# -*- coding: utf-8 -*-
"""Matriz de riesgos DS 44 - Tazki.
Alineada con la GUIA PARA LA IDENTIFICACION Y EVALUACION DE RIESGOS EN LOS
LUGARES DE TRABAJO, version 3, del Instituto de Salud Publica (Res. Ex. E668/25),
que el articulo 7 del DS 44 cita como metodologia rectora.
Metodo VEP (Valor Esperado de la Perdida) = Probabilidad x Consecuencia, con
escalas 1/2/4 y niveles Tolerable / Moderado / Importante / Intolerable.
Formato de columnas segun el Anexo N°6 de la guia; hoja de levantamiento segun Anexo N°2.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, COBALT, LINE = "0C1E50", "1537D1", "C7D0E0"
SOFT, BAND, YELLOW, GREY = "EEF2FE", "F4F6FB", "FFF7D6", "6B76A0"
CALC = "E8EEFB"
FN = "Arial"
def fo(sz=10, b=False, c="0C1E50", it=False):
    return Font(name=FN, size=sz, bold=b, color=c, italic=it)
thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

def banda(ws, rng, texto, alto=30):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = texto; c.font = fo(14, True, "FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[int("".join(x for x in rng.split(":")[0] if x.isdigit()))].height = alto

def sub(ws, rng, texto, alto=22):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = texto; c.font = fo(9, False, "FFFFFF")
    c.fill = PatternFill("solid", fgColor=COBALT)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1, wrap_text=True)
    ws.row_dimensions[int("".join(x for x in rng.split(":")[0] if x.isdigit()))].height = alto

def heads(ws, headers, row, col0=2, alto=46):
    for i, h in enumerate(headers):
        cell = ws.cell(row=row, column=col0+i, value=h)
        cell.font = fo(8.5, True, "FFFFFF")
        cell.fill = PatternFill("solid", fgColor=COBALT)
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = box
    ws.row_dimensions[row].height = alto

def lista(ws, formula, rango):
    dv = DataValidation(type="list", allow_blank=True, formula1=formula)
    ws.add_data_validation(dv); dv.add(rango)

# ══════════════════ HOJA 1 · MATRIZ DE RIESGOS (Anexo 6) ══════════════════
ws = wb.active; ws.title = "Matriz de riesgos"
ws.sheet_view.showGridLines = False
for k, v in [("A",2.5),("B",24),("C",22),("D",24),("E",34),("F",20),("G",20),
             ("H",12),("I",12),("J",9),("K",15),("L",18),("M",15),("N",38),
             ("O",18),("P",12),("Q",13),("R",22)]:
    ws.column_dimensions[k].width = v

banda(ws, "B2:R2", "MATRIZ DE IDENTIFICACIÓN DE PELIGROS/FACTORES DE RIESGO Y EVALUACIÓN DE RIESGOS")
sub(ws, "B3:R3",
    "DS 44, artículo 7 · Formato según Anexo N°6 de la Guía del ISP v3 (Res. Ex. E668/25) · "
    "Método VEP · Complete las celdas amarillas; las azules se calculan solas")

# cabecera de identificacion
ident = [("Empresa","Área"),("Sucursal","Responsable del levantamiento")]
r = 5
for izq, der in ident:
    for col, etiqueta in ((2, izq), (8, der)):
        a = ws.cell(row=r, column=col, value=etiqueta)
        a.font = fo(9, True); a.fill = PatternFill("solid", fgColor=SOFT)
        a.alignment = Alignment(vertical="center", indent=1); a.border = box
        ws.merge_cells(start_row=r, start_column=col+1, end_row=r, end_column=col+5)
        b = ws.cell(row=r, column=col+1)
        b.fill = PatternFill("solid", fgColor=YELLOW)
        for cc in range(col+1, col+6): ws.cell(row=r, column=cc).border = box
    ws.row_dimensions[r].height = 20
    r += 1

HR = 8
H = ["Proceso","Puestos de trabajo","Tareas",
     "Identificación de peligros / factores de riesgo","Riesgo","Tipo de riesgo",
     "Probabilidad (P)","Consecuencia (S)","VEP","Nivel de riesgo (seguridad)",
     "Magnitud de la exposición","Nivel de riesgo (higiénico / psicosocial / MSK)",
     "Medidas preventivas","Responsable","Plazo","Estado","Evidencia de cierre"]
heads(ws, H, HR)

EJ = ["Mantenimiento preventivo de racks de almacenamiento (de apoyo)",
      "Encargado de mantención; técnicos de mantenimiento",
      "Revisión de los racks metálicos de almacenamiento",
      "Fumar en lugares con prohibición de hacerlo; uso de herramientas eléctricas defectuosas; cables y enchufes en mal estado",
      "Incendio / explosión", "Seguridad y emergencias", 4, 2, None, None, None, None,
      "Habilitar zona de fumadores según normativa; verificar estado de herramientas antes de su uso con tarjeta verde/roja; prohibir el uso de herramientas defectuosas",
      "Encargado de mantención", "2026-10-31", "En curso", "Acta de verificación + registro fotográfico"]

FIRST, LAST = HR+1, HR+100
for rr in range(FIRST, LAST+1):
    for i in range(len(H)):
        col = 2+i
        cell = ws.cell(row=rr, column=col)
        cell.border = box; cell.font = fo(9)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if col in (10, 11):                      # J = VEP, K = Nivel seguridad
            cell.fill = PatternFill("solid", fgColor=CALC)
        else:
            cell.fill = PatternFill("solid", fgColor=YELLOW)
    if rr == FIRST:
        for i, v in enumerate(EJ):
            if v is not None:
                ws.cell(row=rr, column=2+i, value=v)
        for i in range(len(H)):
            ws.cell(row=rr, column=2+i).font = fo(9, False, GREY, it=True)
    ws.cell(row=rr, column=16).number_format = "yyyy-mm-dd"   # P = Plazo
    # J (col 10) = VEP, solo para riesgos de seguridad y emergencias
    j = ws.cell(row=rr, column=10)
    j.value = (f'=IF(G{rr}<>"Seguridad y emergencias","",'
               f'IF(OR(H{rr}="",I{rr}=""),"",H{rr}*I{rr}))')
    j.alignment = Alignment(vertical="center", horizontal="center"); j.font = fo(10, True)
    # K (col 11) = Nivel de riesgo segun VEP (tabla de la Guia ISP)
    k = ws.cell(row=rr, column=11)
    k.value = (f'=IF(J{rr}="","",'
               f'IF(J{rr}<=2,"Tolerable",'
               f'IF(J{rr}<=4,"Moderado",'
               f'IF(J{rr}<=8,"Importante","Intolerable"))))')
    k.alignment = Alignment(vertical="center", horizontal="center"); k.font = fo(10, True)

lista(ws, '"Seguridad y emergencias,Higiénico,Psicosocial,Músculo esquelético"', f"G{FIRST}:G{LAST}")
lista(ws, '"1,2,4"', f"H{FIRST}:H{LAST}")
lista(ws, '"1,2,4"', f"I{FIRST}:I{LAST}")
lista(ws, '"Pendiente,En curso,Implementada,Verificada"', f"Q{FIRST}:Q{LAST}")
ws.freeze_panes = f"B{FIRST}"

# pie de firmas (Anexo 6)
PF = LAST + 2
for col, etiqueta in ((2,"Elaborado por"),(6,"Revisado por"),(10,"Aprobado por")):
    a = ws.cell(row=PF, column=col, value=etiqueta)
    a.font = fo(9, True); a.fill = PatternFill("solid", fgColor=SOFT)
    a.alignment = Alignment(vertical="center", indent=1); a.border = box
    for cc in (col+1, col+2):
        ws.cell(row=PF, column=cc).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row=PF, column=cc).border = box
    b = ws.cell(row=PF+1, column=col, value="Fecha")
    b.font = fo(9, True); b.fill = PatternFill("solid", fgColor=SOFT)
    b.alignment = Alignment(vertical="center", indent=1); b.border = box
    for cc in (col+1, col+2):
        ws.cell(row=PF+1, column=cc).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row=PF+1, column=cc).border = box
    ws.cell(row=PF+1, column=col+1).number_format = "yyyy-mm-dd"

# resumen
SR = PF + 3
ws.merge_cells(f"B{SR}:C{SR}")
c = ws[f"B{SR}"]; c.value = "RESUMEN"
c.font = fo(10, True, "FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
res = [("Riesgos evaluados (VEP)", f'=COUNT(J{FIRST}:J{LAST})'),
       ("Intolerables", f'=COUNTIF(K{FIRST}:K{LAST},"Intolerable")'),
       ("Importantes",  f'=COUNTIF(K{FIRST}:K{LAST},"Importante")'),
       ("Moderados",    f'=COUNTIF(K{FIRST}:K{LAST},"Moderado")'),
       ("Tolerables",   f'=COUNTIF(K{FIRST}:K{LAST},"Tolerable")'),
       ("Medidas pendientes", f'=COUNTIF(Q{FIRST}:Q{LAST},"Pendiente")')]
for i,(lab,f_) in enumerate(res):
    rr = SR+1+i
    ws.merge_cells(f"B{rr}:C{rr}")
    a = ws[f"B{rr}"]; a.value = lab; a.font = fo(9, True)
    a.fill = PatternFill("solid", fgColor=SOFT)
    a.alignment = Alignment(vertical="center", indent=1)
    for cc in (2,3): ws.cell(row=rr, column=cc).border = box
    b = ws.cell(row=rr, column=4, value=f_)
    b.font = fo(10, True, COBALT); b.border = box
    b.alignment = Alignment(vertical="center", horizontal="center")

# ══════════════ HOJA 2 · LEVANTAMIENTO DE PROCESOS (Anexo 2) ══════════════
lv = wb.create_sheet("Levantamiento de procesos")
lv.sheet_view.showGridLines = False
for k,v in [("A",2.5),("B",30),("C",26),("D",26),("E",14),("F",22),("G",12),("H",26),("I",24)]:
    lv.column_dimensions[k].width = v
banda(lv, "B2:I2", "LEVANTAMIENTO DE PROCESOS")
sub(lv, "B3:I3",
    "Formato del Anexo N°2 de la Guía del ISP · Es el insumo de la matriz: "
    "primero se levantan procesos, tareas y puestos, y recién después se identifican peligros")
HL = ["Proceso (nombre y tipo: operacional o de apoyo)","Puestos de trabajo involucrados",
      "Tarea","¿Rutinaria?","Lugar específico donde se realiza",
      "N° de personas trabajadoras","Identidad sexogenérica","Observaciones"]
heads(lv, HL, 5, alto=42)
EJL = ["Mantenimiento preventivo de racks de almacenamiento (de apoyo)",
       "Jefe de área; personal de farmacia; encargado de mantención; técnicos de mantenimiento",
       "Revisión de los racks metálicos de almacenamiento","SÍ","Bodegas", 10,
       "1 mujer; 5 mujeres; 1 persona no binaria; 3 hombres","Sin observaciones"]
for rr in range(6, 56):
    for i in range(len(HL)):
        cell = lv.cell(row=rr, column=2+i)
        cell.border = box; cell.font = fo(9)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor=YELLOW)
    if rr == 6:
        for i,v in enumerate(EJL):
            lv.cell(row=rr, column=2+i, value=v).font = fo(9, False, GREY, it=True)
lista(lv, '"SÍ,NO"', "E6:E55")
lv.freeze_panes = "B6"

# ═══════════ HOJA 3 · MEDIDAS DE EMERGENCIA ADICIONALES (art. 7) ═══════════
em = wb.create_sheet("Medidas de emergencia")
em.sheet_view.showGridLines = False
for k,v in [("A",2.5),("B",32),("C",22),("D",18),("E",38),("F",20),("G",14)]:
    em.column_dimensions[k].width = v
banda(em, "B2:G2", "MEDIDAS PREVENTIVAS DE CONTROL Y DE EMERGENCIA ADICIONALES")
sub(em, "B3:G3",
    "El artículo 7 del DS 44 obliga a adoptar medidas preventivas de control y de emergencia "
    "ADICIONALES cuando el riesgo evaluado resulta elevado, alto o grave. Use esta hoja para dejarlas registradas.", alto=32)
HE = ["Situación de emergencia previsible","Proceso o área afectada",
      "Nivel de riesgo asociado","Medidas adicionales de control y emergencia","Responsable","Estado"]
heads(em, HE, 5, alto=36)
EJE = ["Incendio en zona de racks de almacenamiento","Bodegas","Importante",
       "Extintores señalizados y al alcance de todo el personal; detector de humo; procedimiento de evacuación con simulacro semestral",
       "Encargado de mantención","Implementada"]
for rr in range(6, 46):
    for i in range(len(HE)):
        cell = em.cell(row=rr, column=2+i)
        cell.border = box; cell.font = fo(9)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor=YELLOW)
    if rr == 6:
        for i,v in enumerate(EJE):
            em.cell(row=rr, column=2+i, value=v).font = fo(9, False, GREY, it=True)
lista(em, '"Tolerable,Moderado,Importante,Intolerable"', "D6:D45")
lista(em, '"Pendiente,En curso,Implementada,Verificada"', "G6:G45")
em.freeze_panes = "B6"

# ═════════════════ HOJA 4 · CONTROL DE VERSIONES ═════════════════
cv = wb.create_sheet("Control de versiones")
cv.sheet_view.showGridLines = False
for k,v in [("A",2.5),("B",11),("C",14),("D",34),("E",30),("F",22),("G",16),("H",14)]:
    cv.column_dimensions[k].width = v
banda(cv, "B2:H2", "CONTROL DE VERSIONES DE LA MATRIZ")
sub(cv, "B3:H3",
    "La matriz debe revisarse a lo menos anualmente, o cuando cambien las condiciones de trabajo, "
    "ocurra un accidente, se diagnostique una enfermedad profesional o se genere un riesgo grave e inminente. "
    "La Guía del ISP pide mantener los registros históricos de las actualizaciones.", alto=32)
HV = ["Versión","Fecha","Motivo de la actualización","Participantes","Aprobada por",
      "Próxima revisión","Estado"]
heads(cv, HV, 5, alto=30)
EJV = ["v1.0","2026-09-04","Elaboración inicial","Comité paritario y jefaturas de área",
       "Gerencia de operaciones","2027-09-04"]
for rr in range(6, 36):
    for i in range(len(HV)):
        cell = cv.cell(row=rr, column=2+i)
        cell.border = box; cell.font = fo(9)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if i < 6: cell.fill = PatternFill("solid", fgColor=YELLOW)
        else:     cell.fill = PatternFill("solid", fgColor=CALC)
    if rr == 6:
        for i,v in enumerate(EJV):
            cv.cell(row=rr, column=2+i, value=v).font = fo(9, False, GREY, it=True)
    for col in (3, 7):
        cv.cell(row=rr, column=col).number_format = "yyyy-mm-dd"
    e = cv.cell(row=rr, column=8)
    e.value = (f'=IF(C{rr}="","",IF(G{rr}="","Sin próxima revisión",'
               f'IF(G{rr}<TODAY(),"Vencida",IF(G{rr}-TODAY()<=30,"Por vencer","Vigente"))))')
    e.alignment = Alignment(vertical="center", horizontal="center"); e.font = fo(9, True)
cv.freeze_panes = "B6"

# ═════════════════ HOJA 5 · CÓMO USARLA ═════════════════
gu = wb.create_sheet("Cómo usarla")
gu.sheet_view.showGridLines = False
for k,v in [("A",2.5),("B",28),("C",84)]: gu.column_dimensions[k].width = v
banda(gu, "B2:C2", "CÓMO USAR ESTA MATRIZ")
filas = [
 ("De dónde sale este formato",
  "Del Anexo N°6 de la «Guía para la identificación y evaluación de riesgos en los lugares de trabajo», versión 3, del Instituto de Salud Pública, oficializada por Resolución Exenta E668/25. El DS 44 instruye que la matriz de riesgos se elabore en base a las directrices de la guía del ISP vigente, así que esta es la metodología que espera la autoridad."),
 ("El orden de trabajo",
  "Primero se levantan procesos, tareas y puestos (hoja «Levantamiento de procesos», Anexo N°2 de la guía). Recién con eso se identifican los peligros y se evalúan los riesgos en la hoja «Matriz de riesgos». Saltarse el levantamiento es lo que produce matrices genéricas."),
 ("Método VEP",
  "Para los riesgos de seguridad, desastre y emergencias, la guía exige el método del Valor Esperado de la Pérdida: VEP = Probabilidad × Consecuencia. El resultado va de 1 a 16."),
 ("Escala de probabilidad",
  "Solo tres valores: 1 Baja (el daño ocurrirá rara vez, ocurrencia remota) · 2 Media (ocurrirá en varias ocasiones, puede pasar) · 4 Alta (ocurrirá siempre o casi siempre, es evidente que pasará). No existe el valor 3."),
 ("Escala de consecuencia",
  "También tres valores: 1 Ligeramente dañino (cortes superficiales, magulladuras, molestias con recuperación rápida) · 2 Dañino (laceraciones, quemaduras, fracturas pequeñas, intoxicaciones con incapacidad temporal) · 4 Extremadamente dañino (amputaciones, fracturas mayores, incapacidades permanentes, lesiones fatales)."),
 ("Niveles de riesgo",
  "VEP 2 o menos: Tolerable · VEP 4: Moderado · VEP 8: Importante · VEP 16: Intolerable. La planilla los calcula sola."),
 ("Qué hacer con cada nivel",
  "Tolerable: no se necesita mejorar la acción preventiva, pero se requieren comprobaciones periódicas. Moderado: hacer esfuerzos para reducir el riesgo en un plazo determinado. Importante: no comenzar ni continuar el trabajo hasta reducir el riesgo. Intolerable: no comenzar ni continuar; si no es posible reducirlo ni con recursos ilimitados, se debe prohibir el trabajo."),
 ("⚠️ Moderado con consecuencia 4",
  "La propia guía lo advierte: cuando un riesgo Moderado está asociado a consecuencias extremadamente dañinas, se precisa una acción posterior para establecer con más precisión la probabilidad de daño, como base para decidir si hay que mejorar las medidas de control. No lo deje pasar solo porque el número dio 4."),
 ("Los otros tres tipos de riesgo",
  "El VEP aplica a seguridad y emergencias. Los riesgos higiénicos, psicosociales y músculo-esqueléticos se evalúan con la magnitud de la exposición que define su protocolo de vigilancia respectivo (por ejemplo el Protocolo TMERT para MSK), y esa magnitud se anota a mano en sus columnas."),
 ("Enfoque de género",
  "No es una casilla decorativa. La guía pide registrar la identidad sexogenérica de las personas por puesto en el levantamiento, e incorporar los sesgos de género tanto al estimar la probabilidad como al elegir la medida de control. El ejemplo de la guía es claro: un carro de extinción fuera del alcance de las trabajadoras es un hallazgo de la matriz."),
 ("Medidas de control",
  "Aplique la jerarquía de controles: eliminar, sustituir, controles de ingeniería, controles administrativos y por último elementos de protección personal. La guía pide explicitar en la matriz las medidas de aplicación inmediata, y remitir al programa de trabajo preventivo las de mediano y largo plazo."),
 ("Emergencias",
  "El artículo 7 del DS 44 obliga a adoptar medidas preventivas de control y de emergencia adicionales cuando el riesgo evaluado resulta elevado, alto o grave. Para eso está la hoja «Medidas de emergencia»."),
 ("Revisión",
  "A lo menos una vez al año, y además ante cambios de condiciones de trabajo, accidentes, enfermedades profesionales o riesgo grave e inminente. Mantenga los registros históricos en la hoja «Control de versiones»."),
 ("Qué viene después",
  "El artículo 8 del DS 44 da un plazo de 30 días corridos, contados desde que se confecciona o actualiza la matriz, para elaborar o modificar el programa de trabajo preventivo. Y con la matriz lista corresponde difundirla mediante el mapa de riesgos."),
 ("Aviso",
  "Planilla elaborada por Tazki siguiendo la Guía del ISP v3. Es una herramienta de trabajo y no reemplaza la asesoría de un experto en prevención de riesgos ni la asistencia técnica de su organismo administrador de la Ley 16.744."),
]
rr = 4
for tit, txt in filas:
    a = gu.cell(row=rr, column=2, value=tit)
    a.font = fo(10, True, COBALT); a.fill = PatternFill("solid", fgColor=SOFT)
    a.alignment = Alignment(vertical="top", wrap_text=True, indent=1); a.border = box
    b = gu.cell(row=rr, column=3, value=txt)
    b.font = fo(9); b.alignment = Alignment(vertical="top", wrap_text=True); b.border = box
    gu.row_dimensions[rr].height = max(30, 12.5*(len(txt)//86 + 1))
    rr += 1

wb.save("/tmp/piezas/Matriz-IPER-DS44-Tazki.xlsx")
print("guardado")
