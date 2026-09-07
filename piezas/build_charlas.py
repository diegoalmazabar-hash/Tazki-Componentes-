# -*- coding: utf-8 -*-
"""Calendario anual de charlas de 5 minutos + registro de asistencia - Tazki.
Temario construido para que un anio de charlas cubra los contenidos que el
DS 44 exige difundir. Contrastado contra normativa/ds44-texto-completo-literal.md.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

NAVY, COBALT, LINE = "0C1E50", "1537D1", "C7D0E0"
SOFT, BAND, GREY, CALC = "EEF2FE", "F4F6FB", "6B76A0", "E8EEFB"
OK, BAD, NA = "E4F4EC", "FBE9E7", "F0F2F6"
FN = "Arial"
def fo(sz=10, b=False, c="0C1E50", it=False):
    return Font(name=FN, size=sz, bold=b, color=c, italic=it)
thin = Side(style="thin", color=LINE); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

def banda(ws, rng, texto, alto=30):
    ws.merge_cells(rng); c = ws[rng.split(":")[0]]
    c.value = texto; c.font = fo(14, True, "FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[int("".join(x for x in rng.split(":")[0] if x.isdigit()))].height = alto

def sub(ws, rng, texto, alto=24):
    ws.merge_cells(rng); c = ws[rng.split(":")[0]]
    c.value = texto; c.font = fo(9, False, "FFFFFF"); c.fill = PatternFill("solid", fgColor=COBALT)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1, wrap_text=True)
    ws.row_dimensions[int("".join(x for x in rng.split(":")[0] if x.isdigit()))].height = alto

def heads(ws, headers, row, col0=2, alto=42):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col0+i, value=h)
        c.font = fo(8.5, True, "FFFFFF"); c.fill = PatternFill("solid", fgColor=COBALT)
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True); c.border = box
    ws.row_dimensions[row].height = alto

# (tema, de donde sale)
T = [
("Los riesgos de tu puesto: qué dice la matriz sobre lo que haces","Matriz · Art. 7"),
("Qué es el IRL y por qué lo firmaste antes de empezar","Art. 15"),
("Los cuatro contenidos que el IRL debe tener","Art. 15"),
("Orden y aseo: la condición del puesto que más accidentes evita","Art. 15 N°1"),
("Cómo se eligen las medidas: primero eliminar, al final el EPP","Art. 9"),
("Tu EPP: colocación, limitaciones de uso y chequeo diario","Art. 13"),
("Cuándo se cambia un EPP y a quién se le pide","Art. 13"),
("Protección colectiva: por qué la baranda vale más que el arnés","Art. 12"),
("Máquinas y herramientas: atrapamiento, corte y amputación","Art. 10"),
("Las protecciones de la máquina no se retiran: qué hacer si falta una","Art. 10"),
("Qué dicen los manuales y fichas técnicas de tus equipos","Art. 10"),
("Riesgo grave e inminente: tienes derecho a parar","Art. 18"),
("Qué pasa después de que paras: aviso, evaluación y reanudación","Art. 18"),
("El plan de emergencia de nuestro lugar de trabajo","Art. 19"),
("Vías de evacuación y zona de seguridad: recorrido real","Art. 19"),
("El simulacro anual: para qué sirve y qué se evalúa","Art. 19"),
("Señalética: qué significa cada color y cada forma","Art. 16 N°6"),
("Prevención de incendios y uso del extintor","Art. 16 N°7"),
("El mapa de riesgos de tu área: dónde está y cómo se lee","Art. 62"),
("Qué hacer ante un accidente: primeros pasos y a quién avisar","Art. 16 N°4"),
("Tus prestaciones de la ley 16.744 y a qué centro asistencial ir","Art. 16 N°4"),
("Accidentes de trayecto: el camino también cuenta","Art. 4 N°4"),
("Efectos en la salud de los riesgos a los que estás expuesto","Art. 16 N°2"),
("Enfermedades profesionales: cómo se reconocen y qué hacer","Art. 16 N°2"),
("Vigilancia de la salud: por qué te citan a exámenes","Art. 67"),
("El tiempo del examen es tiempo trabajado","Art. 68"),
("Riesgos ergonómicos: posturas, carga y repetición","Art. 7"),
("Manejo manual de carga: técnica y límites","Art. 7"),
("Riesgos psicosociales: qué son y cómo se miden","Art. 7"),
("Violencia y acoso en el trabajo: qué es y dónde denunciar","Art. 7"),
("Personas especialmente sensibles: cuándo avisar tu condición","Art. 11"),
("Sustancias peligrosas: etiqueta, ficha de seguridad y almacenamiento","Art. 15 N°4"),
("Procedimientos de trabajo seguro: por qué existen","Art. 15 N°3"),
("Trabajos en altura","Matriz · Art. 7"),
("Espacios confinados","Matriz · Art. 7"),
("Riesgo eléctrico y bloqueo de energías","Matriz · Art. 7"),
("Tránsito de vehículos y peatones en el recinto","Art. 8"),
("Conducción de vehículos: fatiga, distracción y velocidad","Art. 8"),
("Alcohol y drogas en el trabajo: por qué es parte del programa","Art. 8"),
("Alimentación y vida saludable: el otro lado de la prevención","Art. 8"),
("Qué es el Comité Paritario y quiénes lo integran","Arts. 23 y 47"),
("Qué es el Delegado de Seguridad y Salud y cómo se elige","Art. 66"),
("Cómo hacer una propuesta o un reclamo de seguridad","Art. 58 N°2 h)"),
("Investigación de accidentes: por qué se pregunta tanto","Art. 71"),
("Incidentes sin lesión: por qué también se reportan","Art. 73 N°1"),
("El programa de trabajo preventivo: qué comprometió la empresa","Art. 8"),
("El reglamento interno: dónde está y qué te obliga","Arts. 56 y 59"),
("Prohibiciones del reglamento interno y sus sanciones","Arts. 60 y 61"),
("Contratistas y visitas: cómo se coordinan los riesgos","Art. 20"),
("Repaso del año: los riesgos que más se repitieron","Registro propio"),
("Repaso del año: qué medidas cerramos y cuáles quedaron abiertas","Art. 14"),
("Cierre y compromisos para el próximo período","Art. 14"),
]

# ══════════ HOJA 1 · CALENDARIO ANUAL ══════════
ws = wb.active; ws.title = "Calendario anual"
ws.sheet_view.showGridLines = False
for k,v in [("A",2.5),("B",7),("C",13),("D",58),("E",20),("F",22),("G",13),("H",14),("I",11),("J",30)]:
    ws.column_dimensions[k].width = v
banda(ws,"B2:J2","CALENDARIO ANUAL DE CHARLAS DE 5 MINUTOS")
sub(ws,"B3:J3","52 temas para que un año de charlas cubra lo que el DS 44 obliga a difundir. "
               "Cambia los que no apliquen a tu operación y agrega los riesgos propios de tu matriz. "
               "La columna «De dónde sale» indica el artículo que respalda el tema.")
heads(ws,["N°","Fecha planificada","Tema de la charla","De dónde sale","Relator o responsable",
          "Estado","Fecha realizada","Asistentes","Observaciones"],5)
r0=6
for i,(tema,src) in enumerate(T):
    r=r0+i
    for j,v in enumerate([i+1,"",tema,src,"","","","",""]):
        c=ws.cell(row=r,column=2+j,value=v); c.border=box; c.font=fo(9)
        c.alignment=Alignment(vertical="top",wrap_text=True,
                              horizontal="center" if j in (0,1,5,6,7) else "left")
        if j in (1,4,5,6,7,8): c.fill=PatternFill("solid",fgColor=BAND)
    ws.cell(row=r,column=3).number_format="dd-mm-yyyy"
    ws.cell(row=r,column=8).number_format="dd-mm-yyyy"
    ws.row_dimensions[r].height=28
rN=r0+len(T)-1
dv=DataValidation(type="list",allow_blank=True,formula1='"Planificada,Realizada,Reprogramada,No aplica"')
ws.add_data_validation(dv); dv.add("G%d:G%d"%(r0,rN))
rng="G%d:G%d"%(r0,rN)
ws.conditional_formatting.add(rng,FormulaRule(formula=['$G%d="Realizada"'%r0],fill=PatternFill("solid",fgColor=OK)))
ws.conditional_formatting.add(rng,FormulaRule(formula=['$G%d="Reprogramada"'%r0],fill=PatternFill("solid",fgColor=BAD)))
ws.conditional_formatting.add(rng,FormulaRule(formula=['$G%d="No aplica"'%r0],fill=PatternFill("solid",fgColor=NA)))
ws.freeze_panes="D6"; ws.auto_filter.ref="B5:J%d"%rN

# ══════════ HOJA 2 · REGISTRO DE ASISTENCIA ══════════
ws2=wb.create_sheet("Registro de asistencia")
ws2.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",13),("C",44),("D",22),("E",20),("F",26),("G",14),("H",22),("I",18),("J",26)]:
    ws2.column_dimensions[k].width=v
banda(ws2,"B2:J2","REGISTRO DE ASISTENCIA A CHARLAS")
sub(ws2,"B3:J3","Una fila por persona y por charla. Es la evidencia que se pide en una fiscalización: "
                "sin registro firmado, la charla no ocurrió. Imprime la hoja para la firma o usa firma digital.")
heads(ws2,["Fecha","Tema de la charla","Relator","Área o lugar de trabajo","Nombre de la persona",
           "RUT","Cargo","Firma","Observaciones"],5,alto=36)
for i in range(300):
    r=6+i
    for col in range(2,11):
        c=ws2.cell(row=r,column=col); c.border=box; c.font=fo(9)
        c.alignment=Alignment(vertical="center",wrap_text=(col in (3,10)),
                              horizontal="center" if col in (2,7) else "left")
    ws2.cell(row=r,column=2).number_format="dd-mm-yyyy"
    ws2.row_dimensions[r].height=22
ws2.freeze_panes="B6"; ws2.auto_filter.ref="B5:J305"

# ══════════ HOJA 3 · RESUMEN ══════════
ws3=wb.create_sheet("Resumen")
ws3.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",42),("C",16),("D",68)]:
    ws3.column_dimensions[k].width=v
banda(ws3,"B2:D2","RESUMEN DEL AÑO")
sub(ws3,"B3:D3","Se calcula solo a partir de las dos hojas anteriores.")
FIL=[("Charlas planificadas en el calendario",'=COUNTA(\'Calendario anual\'!$D$%d:$D$%d)'%(r0,rN),"Temas cargados en el calendario."),
     ("Charlas realizadas",'=COUNTIF(\'Calendario anual\'!$G$%d:$G$%d,"Realizada")'%(r0,rN),"Las marcadas como Realizada."),
     ("Charlas reprogramadas",'=COUNTIF(\'Calendario anual\'!$G$%d:$G$%d,"Reprogramada")'%(r0,rN),"Ojo si son muchas: es la señal de que el calendario no se está sosteniendo."),
     ("Charlas marcadas No aplica",'=COUNTIF(\'Calendario anual\'!$G$%d:$G$%d,"No aplica")'%(r0,rN),"Quedan fuera del porcentaje de cumplimiento."),
     ("% de cumplimiento del calendario",'=IF((C7+C8)=0,"—",C7/(C7+C8))',"Realizadas sobre realizadas más reprogramadas."),
     ("Registros de asistencia cargados",'=COUNTA(\'Registro de asistencia\'!$F$6:$F$305)',"Total de firmas registradas en el año."),
     ("Charlas distintas con asistencia registrada",'=SUMPRODUCT((\'Registro de asistencia\'!$C$6:$C$305<>"")/COUNTIF(\'Registro de asistencia\'!$C$6:$C$305,\'Registro de asistencia\'!$C$6:$C$305&""))',"Debería acercarse al número de charlas realizadas. Si es menor, hay charlas hechas sin registro."),
     ("Promedio de asistentes por charla",'=IF(C12=0,"—",C11/C12)',"Referencia para detectar charlas con baja convocatoria.")]
r=6
for nom,frm,nota in FIL:
    a=ws3.cell(row=r,column=2,value=nom); a.font=fo(10,True); a.border=box
    a.alignment=Alignment(vertical="center",wrap_text=True)
    b=ws3.cell(row=r,column=3,value=frm); b.font=fo(10,True); b.border=box
    b.fill=PatternFill("solid",fgColor=CALC); b.alignment=Alignment(horizontal="center",vertical="center")
    c=ws3.cell(row=r,column=4,value=nota); c.font=fo(9,False,GREY); c.border=box
    c.alignment=Alignment(vertical="center",wrap_text=True)
    ws3.row_dimensions[r].height=32
    r+=1
ws3["C10"].number_format="0%"; ws3["C13"].number_format="0.0"
ws3.cell(row=r+1,column=2,value="Nota: el % de cumplimiento excluye las charlas marcadas «No aplica».").font=fo(8.5,False,GREY)

# ══════════ HOJA 4 · CÓMO USARLA ══════════
ws4=wb.create_sheet("Cómo usarla")
ws4.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",28),("C",104)]: ws4.column_dimensions[k].width=v
banda(ws4,"B2:C2","CÓMO USAR ESTE CALENDARIO")
F=[("Lo primero, y es importante","El DS 44 no menciona la «charla de 5 minutos» con ese nombre. La única vez que el decreto dice «charlas» es en el artículo 47, entre los medios con que el Comité Paritario cumple su función de instruir a las personas trabajadoras: «organizando reuniones informativas, charlas o cualquier otro medio de divulgación»."),
   ("Entonces, ¿es obligatoria?","La charla en sí no. Lo que sí es obligatorio es aquello que la charla sirve para cumplir: las acciones permanentes de difusión y promoción del artículo 4 N°4, la información de riesgos del artículo 15 cuando cambian los procesos, los materiales o las sustancias, y la función de divulgación del Comité del artículo 47. La charla diaria es el vehículo más práctico para sostener eso todo el año, y es lo que un fiscalizador espera ver documentado."),
   ("Lo que la charla NO reemplaza","La capacitación del artículo 16, que es un curso de al menos 8 horas con siete temas definidos y una periodicidad que no puede superar los dos años. Y la capacitación de EPP del artículo 13, de al menos una hora cronológica, reforzada cada año. Sumar charlas de cinco minutos no equivale a esas horas."),
   ("Cómo se completa","Ponle fecha a cada tema al inicio del período. Marca el estado a medida que avanzas. Y por cada charla realizada, carga las firmas en la hoja «Registro de asistencia»: una fila por persona."),
   ("Por qué el temario está armado así","Los 52 temas recorren los contenidos que el decreto obliga a informar y capacitar. La columna «De dónde sale» indica el artículo de respaldo. Reemplaza los que no apliquen a tu operación por los riesgos propios de tu matriz: los temas marcados «Matriz · Art. 7» están ahí justamente para eso."),
   ("La evidencia que se pide","Registro con fecha, tema, relator y firma de los asistentes. Sin eso, para la fiscalización la charla no ocurrió: el artículo 72 exige que toda la gestión esté registrada y respaldada de forma documental y fidedigna."),
   ("Un error frecuente","Registrar siempre a las mismas personas. Si en tu área hay turnos, la misma charla se repite por turno y cada uno tiene su registro. La hoja de asistencia admite repetir el tema en fechas distintas."),
   ("Quién lo hizo","Tazki, software chileno de gestión de seguridad y salud en el trabajo. Si quieres que estos registros dejen de vivir en carpetas y planillas, escríbenos en tazki.cl")]
r=4
for t,d in F:
    a=ws4.cell(row=r,column=2,value=t); a.font=fo(10,True); a.fill=PatternFill("solid",fgColor=SOFT)
    a.alignment=Alignment(vertical="top",wrap_text=True); a.border=box
    b=ws4.cell(row=r,column=3,value=d); b.font=fo(10); b.alignment=Alignment(vertical="top",wrap_text=True); b.border=box
    ws4.row_dimensions[r].height=64; r+=1

# ══════════ HOJA 5 · CONTROL DE VERSIONES ══════════
ws5=wb.create_sheet("Control de versiones")
ws5.sheet_view.showGridLines=False
for k,v in [("A",2.5),("B",14),("C",16),("D",30),("E",56)]: ws5.column_dimensions[k].width=v
banda(ws5,"B2:E2","CONTROL DE VERSIONES")
heads(ws5,["Versión","Fecha","Responsable","Qué cambió"],4,alto=26)
for i in range(15):
    r=5+i
    for col in range(2,6):
        c=ws5.cell(row=r,column=col); c.border=box; c.font=fo(9)
        c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="center" if col in (2,3) else "left")
    ws5.cell(row=r,column=3).number_format="dd-mm-yyyy"; ws5.row_dimensions[r].height=24

for s in wb.worksheets: s.sheet_properties.tabColor=COBALT
wb.save("/tmp/piezas/Calendario-Charlas-5min-DS44-Tazki.xlsx")
print("guardado ·", len(T), "temas · hojas:", [s.title for s in wb.worksheets], "· rango calendario", r0, rN)
