# -*- coding: utf-8 -*-
"""Checklist de fiscalizacion DS 44 - Tazki.
Cada linea cita el articulo del decreto que la respalda. Texto contrastado
contra normativa/ds44-texto-completo-literal.md (transcripcion literal de Ley
Chile, DO 27-07-2024, vigente desde 01-02-2025).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

NAVY, COBALT, LINE = "0C1E50", "1537D1", "C7D0E0"
SOFT, BAND, YELLOW, GREY = "EEF2FE", "F4F6FB", "FFF7D6", "6B76A0"
CALC = "E8EEFB"
OK, BAD, NA = "E4F4EC", "FBE9E7", "F0F2F6"
FN = "Arial"
def fo(sz=10, b=False, c="0C1E50", it=False):
    return Font(name=FN, size=sz, bold=b, color=c, italic=it)
thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

def banda(ws, rng, texto, alto=30):
    ws.merge_cells(rng); c = ws[rng.split(":")[0]]
    c.value = texto; c.font = fo(14, True, "FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[int("".join(x for x in rng.split(":")[0] if x.isdigit()))].height = alto

def sub(ws, rng, texto, alto=22):
    ws.merge_cells(rng); c = ws[rng.split(":")[0]]
    c.value = texto; c.font = fo(9, False, "FFFFFF")
    c.fill = PatternFill("solid", fgColor=COBALT)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1, wrap_text=True)
    ws.row_dimensions[int("".join(x for x in rng.split(":")[0] if x.isdigit()))].height = alto

def heads(ws, headers, row, col0=2, alto=44):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col0+i, value=h)
        c.font = fo(8.5, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=COBALT)
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = box
    ws.row_dimensions[row].height = alto

# ══════════════ FILAS DEL CHECKLIST ══════════════
# (ambito, que se verifica, articulo, aplica a, evidencia)
R = [
("Gestión general","Existe una matriz de identificación de peligros y evaluación de riesgos y un programa de gestión de riesgos implementados en los lugares de trabajo, con enfoque de género, mejora continua y participación de las personas trabajadoras.","Art. 4 N°2","Todas","Matriz y programa vigentes, con registro de la participación"),
("Gestión general","Se realizan acciones permanentes de difusión y promoción de la seguridad y salud, incluida la prevención de los riesgos de traslado y de los accidentes de trayecto, con enfoque de género.","Art. 4 N°4","Todas","Campañas, charlas, afiches, correos, registros de difusión"),
("Gestión general","Se permite el ingreso a los lugares de trabajo a las entidades fiscalizadoras y al organismo administrador.","Art. 4 N°3","Todas","Procedimiento de acceso y registro de visitas"),
("Gestión general","Existe un Sistema de Gestión de SST con sus cinco elementos, cada vez que la normativa lo establece.","Art. 22","Cuando la normativa lo exija","Política, estructura organizacional, diagnóstico y planificación, evaluación o auditoría periódica, y acciones de mejora continua"),
("Gestión general","El programa del Sistema de Gestión está aprobado por el representante legal y fue dado a conocer al Comité Paritario y al Departamento de Prevención.","Art. 22 N°3","Con SG-SST","Documento aprobado con fecha y registro de difusión"),
("Matriz (MIPER)","La matriz cubre todos los procesos, tareas y puestos de trabajo.","Art. 7","Todas","Inventario de procesos y puestos contrastado con la matriz"),
("Matriz (MIPER)","Considera la exposición a agentes y factores de riesgo: ergonómicos, psicosociales, violencia y acoso, accidentes y enfermedades ocurridos y riesgos de los programas de vigilancia, con enfoque de género.","Art. 7","Todas","Matriz con esas categorías presentes y trazables"),
("Matriz (MIPER)","Contiene los dos elementos mínimos: identificación de los peligros del puesto de trabajo y evaluación de los riesgos.","Art. 7","Todas","Matriz con ambas columnas completas"),
("Matriz (MIPER)","La evaluación usa una metodología validada y basada en criterios definidos por la autoridad competente, rigiéndose por la Guía Técnica del Instituto de Salud Pública.","Art. 7","Todas","Metodología declarada en la matriz y coherente con la guía del ISP"),
("Matriz (MIPER)","Considera las condiciones de trabajo previsibles a futuro y la situación de las personas especialmente sensibles.","Arts. 7 y 11","Todas","Registro de personas sensibles y medidas específicas adoptadas"),
("Matriz (MIPER)","Se adoptaron medidas preventivas de control y de emergencia adicionales donde el riesgo resultó elevado, alto o grave.","Art. 7","Todas","Medidas adicionales asociadas a esos riesgos en el programa"),
("Matriz (MIPER)","Está disponible en los lugares de trabajo y fue informada a las personas trabajadoras, al Comité Paritario, al Delegado de SST y a los dirigentes sindicales.","Art. 7","Todas","Actas o registros de entrega e información"),
("Matriz (MIPER)","Es conocida por toda la línea de mando.","Art. 7","Todas","Registro de difusión a jefaturas y supervisores"),
("Matriz (MIPER)","Se revisa al menos anualmente y cada vez que cambian las condiciones de trabajo, ocurre un accidente, se diagnostica una enfermedad profesional o se genera una situación de riesgo grave e inminente.","Art. 7","Todas","Historial de versiones con fecha y motivo de cada revisión"),
("Programa preventivo","Se elaboró o modificó dentro de los 30 días corridos contados desde la confección o actualización de la matriz.","Art. 8","Todas","Fechas de la matriz y del programa comparables entre sí"),
("Programa preventivo","Contiene, al menos, las medidas preventivas y correctivas a implementar, los plazos de implementación y los responsables de su ejecución.","Art. 8","Todas","Programa con esas tres columnas completas"),
("Programa preventivo","Incluye actividades de promoción para prevenir los riesgos asociados al consumo de alcohol y drogas y para difundir un estilo de vida y alimentación saludables.","Art. 8","Todas","Actividades programadas y sus registros de ejecución"),
("Programa preventivo","Contempla las actividades para prevenir los riesgos asociados a la conducción de vehículos motorizados, cuando corresponde.","Art. 8","Cuando corresponda","Actividades programadas para conductores"),
("Programa preventivo","Consta por escrito y fue aprobado por el representante legal, indicando la fecha de aprobación y la de sus modificaciones.","Art. 8","Todas","Programa firmado con fecha de aprobación visible"),
("Programa preventivo","Fue difundido antes de su implementación mediante avisos o informaciones visibles en los lugares de trabajo o por correo electrónico, y se remitió un ejemplar al Comité Paritario.","Art. 8","Todas","Evidencia de difusión previa y constancia de envío al Comité"),
("Programa preventivo","Define acciones para controlar y vigilar el cumplimiento de las medidas, con la periodicidad y en los casos que el propio programa establece, presenciales o por medios electrónicos idóneos.","Art. 8","Todas","Programa con el mecanismo de control descrito y sus registros"),
("Programa preventivo","Se realiza, al menos anualmente, una evaluación del cumplimiento del programa que se pronuncia sobre la eficacia de las acciones y dispone las mejoras necesarias.","Art. 14","Todas","Informe anual de evaluación de eficacia con medidas de mejora"),
("Control de riesgos y EPP","Las medidas siguen el orden de prelación: evitar o eliminar, controlar en la fuente, reducir por medidas organizacionales y, solo entonces, EPP.","Art. 9","Todas","Justificación del control elegido para cada riesgo"),
("Control de riesgos y EPP","Se privilegian los mecanismos o equipos de protección colectiva por sobre los elementos de protección personal.","Art. 12","Todas","Registro de las medidas colectivas implementadas"),
("Control de riesgos y EPP","Existe un procedimiento de trabajo seguro para máquinas, equipos y herramientas motrices con riesgo de atrapamiento, corte, lesión o amputación, con programa de operación y mantenimiento, control permanente y protecciones.","Art. 10","Con esas máquinas","Procedimiento escrito, plan de mantenimiento y registro de protecciones"),
("Control de riesgos y EPP","Se informa a las personas trabajadoras sobre el contenido sustancial de los manuales, instrucciones y fichas técnicas de máquinas y equipos.","Art. 10","Con máquinas y equipos","Registro de la información entregada"),
("Control de riesgos y EPP","Los elementos de protección personal son adecuados al riesgo y se proveen a costo de la entidad empleadora.","Art. 13","Todas","Registro de entrega firmado y respaldo de compra"),
("Control de riesgos y EPP","Existe un procedimiento de utilización, mantenimiento, reposición y recambio de los EPP.","Art. 13","Todas","Procedimiento escrito y registros de reposición"),
("Control de riesgos y EPP","Los EPP cumplen las normas vigentes de certificación de calidad o están registrados en el Instituto de Salud Pública.","Art. 13","Todas","Certificados o número de registro ISP de cada elemento"),
("Control de riesgos y EPP","Existe un programa de capacitación sobre uso y mantención de EPP de al menos una hora cronológica, reforzado anualmente y cada vez que ingresa una persona o cambia el tipo de EPP.","Art. 13","Todas","Programa de capacitación y registros de cada sesión"),
("Control de riesgos y EPP","Esas capacitaciones se registran indicando actividades teóricas y prácticas, asistentes, relatores, resultados de las evaluaciones de aprendizaje y actividades de reforzamiento.","Art. 13","Todas","Registro con esos cinco datos"),
("Control de riesgos y EPP","Se identifican las personas especialmente sensibles y no se les asigna a puestos donde esa sensibilidad implique un riesgo grave para su vida o salud o la de terceros.","Art. 11","Todas","Registro reservado de casos y medidas específicas"),
("Información y capacitación","Cada persona trabajadora recibe la información de los riesgos laborales (IRL) de forma oportuna y adecuada, previo al inicio de sus labores.","Art. 15","Todas","IRL firmado con fecha anterior al inicio de labores"),
("Información y capacitación","Se informa nuevamente cada vez que la persona se incorpora a un nuevo proceso productivo o cambian las tecnologías, los materiales o las sustancias utilizadas.","Art. 15","Todas","Registros de re-entrega asociados al cambio"),
("Información y capacitación","El IRL contiene los cuatro contenidos mínimos: características del lugar de trabajo, riesgos y medidas preventivas incluidas las de emergencia, procedimientos de trabajo seguro y características de productos y sustancias.","Art. 15","Todas","Formato de IRL con los cuatro contenidos"),
("Información y capacitación","Mientras están pendientes la matriz y el programa, igual se informan los riesgos inherentes a la actividad.","Art. 15","Todas","IRL provisorio con los riesgos de la actividad"),
("Información y capacitación","Se efectúa la capacitación en prevención con la periodicidad que define el programa, sin exceder dos años.","Art. 16","Todas","Registro con la fecha de la última capacitación de cada persona"),
("Información y capacitación","El curso dura al menos 8 horas y aborda los siete temas del artículo 16, con enfoque de género.","Art. 16","Todas","Programa del curso, control de asistencia y horas"),
("Información y capacitación","Se emplean metodologías que procuren un adecuado aprendizaje y se verifica la comprensión.","Art. 16","Todas","Evaluaciones de aprendizaje aplicadas"),
("Participación","Se promueve la consulta y participación de los representantes de las personas trabajadoras cuando se prevén cambios en los procesos o en la estructura organizacional que puedan poner en riesgo grave la vida y salud.","Art. 17","Todas","Actas de consulta previas al cambio"),
("Participación","Funciona un Comité Paritario de Higiene y Seguridad en toda empresa, faena, sucursal o agencia en que trabajen más de 25 personas.","Art. 23","Más de 25 por faena","Acta de constitución vigente"),
("Participación","El Comité está compuesto por tres representantes de la entidad empleadora y tres de las personas trabajadoras, con un suplente por cada titular.","Art. 25","Con Comité","Nómina de titulares y suplentes"),
("Participación","El acta de constitución fue registrada en el sitio web de la Dirección del Trabajo dentro de los 15 días hábiles siguientes a la elección.","Art. 36","Con Comité","Comprobante de registro con fecha"),
("Participación","El Comité se reúne en forma ordinaria una vez al mes.","Art. 39","Con Comité","Actas de las últimas reuniones mensuales"),
("Participación","La entidad empleadora otorga las facilidades y adopta las medidas necesarias para que el Comité funcione adecuadamente.","Art. 37","Con Comité","Registro de horas, espacio y recursos asignados"),
("Participación","Se eligió un Delegado de Seguridad y Salud en el Trabajo en todo lugar de trabajo o faena donde laboran entre 10 y hasta 25 personas y no funciona un Comité Paritario.","Art. 66","Entre 10 y 25 por faena","Acta de la asamblea de elección"),
("Participación","El acta de la elección del Delegado indica si fue presencial o por medios electrónicos idóneos, y el mandato de hasta dos años está vigente.","Arts. 33 y 66","Con Delegado","Acta con esa constancia y fecha de elección"),
("Riesgo grave y emergencias","Existe procedimiento para informar de inmediato a las personas afectadas cuando sobreviene un riesgo grave e inminente, junto con las medidas adoptadas para eliminarlo o atenuarlo.","Art. 18","Todas","Procedimiento escrito y registros de aplicación"),
("Riesgo grave y emergencias","Existe procedimiento de suspensión inmediata de las faenas afectadas y de evacuación cuando el riesgo no se puede eliminar o atenuar.","Art. 18","Todas","Procedimiento escrito"),
("Riesgo grave y emergencias","Cuando una persona interrumpe sus labores por riesgo grave e inminente, se informa la suspensión a la Inspección del Trabajo respectiva.","Art. 18","Todas","Constancia del aviso a la Inspección"),
("Riesgo grave y emergencias","Existe uno o más planes de gestión, reducción y respuesta de riesgos para emergencias, catástrofes o desastres, internos o externos, conocidos, probables y previsibles.","Art. 19","Todas","Plan documentado y difundido"),
("Riesgo grave y emergencias","El plan se ensaya a lo menos una vez al año, simulando una emergencia real.","Art. 19","Todas","Acta del simulacro con fecha y participantes"),
("Coordinación","Cuando en el lugar de trabajo prestan servicios dos o más entidades empleadoras, o una entidad y personas independientes, todas coordinan y cooperan para la aplicación de las medidas de seguridad.","Art. 20","Con concurrencia","Actas de coordinación e intercambio de información de riesgos"),
("Estructura preventiva","Existe Departamento de Prevención de Riesgos dirigido por un experto, con autonomía en las materias técnicas, en las entidades con más de 100 trabajadores.","Art. 50","Más de 100","Estructura, nombramiento y registro del experto"),
("Estructura preventiva","El experto a cargo tiene la categoría y el tiempo de dedicación mínima que corresponden a la actividad y al tamaño de la entidad.","Arts. 53, 54 y 55","Con Departamento","Título, registro y contrato con dedicación acreditada"),
("Estructura preventiva","En entidades de hasta cien personas, el representante legal solicitó al organismo administrador que lo capaciten, a él o a quien designe, en materias de gestión de riesgos.","Art. 65","Hasta 100","Solicitud y certificado de la capacitación"),
("Estructura preventiva","En entidades de hasta veinticinco personas se aplican los elementos mínimos del artículo 64: política, autoevaluación, matriz y programa según la pauta del organismo administrador.","Art. 64","Hasta 25","Instrumento de autoevaluación aplicado, matriz y programa"),
("Reglamento interno","Existe un Reglamento Interno de Higiene y Seguridad en el Trabajo, mantenido al día, y se entrega gratuitamente un ejemplar a cada persona trabajadora.","Art. 56","Todas","Reglamento vigente y registro de entrega"),
("Reglamento interno","El reglamento fue ingresado en las páginas web de la Seremi de Salud y de la Dirección del Trabajo.","Art. 57","Todas","Comprobantes de ingreso en ambas instituciones"),
("Reglamento interno","Fue remitido al Comité Paritario o al Delegado, a las organizaciones sindicales y a las personas trabajadoras con una anticipación no inferior a 30 días antes de comenzar a regir.","Art. 57","Todas","Constancia de la remisión con fecha"),
("Reglamento interno","Las observaciones se recibieron por escrito dentro de los 10 días corridos siguientes a su comunicación y las aceptadas se incorporaron al texto.","Art. 57","Todas","Observaciones recibidas y versión modificada"),
("Reglamento interno","Se revisa con una periodicidad no inferior a un año, con participación del Departamento de Prevención, del Comité Paritario o del Delegado.","Art. 57","Todas","Acta de la revisión anual"),
("Reglamento interno","Contiene el preámbulo y el capítulo de disposiciones generales con todos sus contenidos mínimos.","Art. 58","Todas","Índice del reglamento contrastado con el artículo 58"),
("Mapas de riesgos","Se mantienen mapas de riesgos en las dependencias, que permiten localizar y visualizar los principales riesgos.","Art. 62","Todas","Mapas impresos o publicados por lugar de trabajo"),
("Mapas de riesgos","Cada mapa contiene, al menos, un dibujo o esquema del lugar de trabajo y los principales riesgos indicados con símbolos, conforme fueron determinados en la matriz.","Art. 62","Todas","Mapa con esos dos elementos y coherencia con la matriz"),
("Mapas de riesgos","El mapa se actualiza cada vez que cambia la matriz de identificación de peligros y evaluación de riesgos.","Art. 62","Todas","Versión y fecha del mapa comparadas con las de la matriz"),
("Mapas de riesgos","Los mapas están disponibles para las personas trabajadoras en sitios visibles de cada lugar de trabajo.","Art. 62","Todas","Registro fotográfico de su ubicación"),
("Mapas de riesgos","Fueron elaborados con la participación del Comité Paritario o, si no existe, del Delegado de Seguridad y Salud en el Trabajo.","Art. 63","Todas","Acta de participación en la confección"),
("Mapas de riesgos","Las actualizaciones siguen la metodología de la guía dictada por el Ministerio del Trabajo y Previsión Social mediante resolución.","Art. 62 · Res. Ex. N°129 de 2025","Todas","Mapa actualizado conforme a la guía vigente desde el 1-3-2025"),
("Vigilancia de la salud","Evaluado el agente de riesgo que pueda causar una enfermedad profesional, se solicitó la incorporación al programa de vigilancia ambiental y de la salud y este se ejecuta.","Art. 67","Con agentes de riesgo","Solicitud al organismo administrador y registros del programa"),
("Vigilancia de la salud","Se guarda reserva de los datos sensibles relacionados con el estado de salud de las personas trabajadoras.","Art. 67","Todas","Procedimiento de resguardo y acceso restringido"),
("Vigilancia de la salud","Se autoriza la asistencia a los exámenes de control y ese tiempo se considera trabajado para todos los efectos legales.","Art. 68","Todas","Autorizaciones emitidas y registro de asistencia"),
("Vigilancia de la salud","Se traslada a la persona con enfermedad profesional a un puesto sin exposición al riesgo que la originó, sin detrimento de sus remuneraciones.","Art. 69","Cuando corresponda","Prescripción del organismo administrador y constancia del traslado"),
("Investigación","Se investigan las causas de los accidentes del trabajo, incidentes peligrosos y enfermedades profesionales, con enfoque de género y promoviendo la participación de las personas trabajadoras y sus representantes.","Art. 71","Todas","Informes de investigación con medidas correctivas"),
("Investigación","Se emplea la metodología de investigación que indica el organismo administrador.","Art. 71","Todas","Metodología declarada en los informes"),
("Registros e indicadores","Toda la información de la gestión de riesgos está registrada y respaldada de forma documental y fidedigna, preferentemente electrónica, a disposición de la entidad fiscalizadora y del organismo administrador.","Art. 72","Todas","Repositorio único, ordenado y accesible"),
("Registros e indicadores","Con Departamento de Prevención: se llevan los registros de incidentes o sucesos peligrosos, de accidentes del trabajo, de trayecto y enfermedades profesionales, y de las personas en vigilancia de la salud.","Art. 73","Con Departamento","Los tres registros con sus datos mínimos"),
("Registros e indicadores","Con Departamento de Prevención: se determinan y mantienen los indicadores de accidentabilidad, frecuencia y gravedad.","Art. 73","Con Departamento","Planilla o tablero de indicadores con su periodicidad"),
("Registros e indicadores","Sin Departamento de Prevención: se registra al menos la tasa anual de accidentabilidad y todos los accidentes del trabajo, de trayecto y enfermedades profesionales, con nombre, sexo, lugar, descripción y relato.","Art. 75","Sin Departamento","Registro con esos datos mínimos"),
("Registros e indicadores","Los registros están diferenciados por sexo respecto de la exposición a factores de riesgo, los accidentes y enfermedades y las personas en vigilancia de la salud.","Art. 74","Todas","Registros con la columna de sexo poblada"),
]

# ══════════════ HOJA 1 · CHECKLIST ══════════════
ws = wb.active; ws.title = "Checklist"
ws.sheet_view.showGridLines = False
for k, v in [("A",2.5),("B",6),("C",20),("D",62),("E",17),("F",17),("G",40),
             ("H",14),("I",18),("J",14),("K",34)]:
    ws.column_dimensions[k].width = v
banda(ws, "B2:K2", "CHECKLIST DE FISCALIZACIÓN · DECRETO SUPREMO N° 44")
sub(ws, "B3:K3", "Preauditoría de cumplimiento. Cada línea indica el artículo del DS 44 que la respalda. "
                 "Marca el estado en la columna «Estado» y las brechas se cuentan solas en la hoja «Resumen».")
HEAD = ["N°","Ámbito","Qué se verifica","Artículo DS 44","A quién aplica",
        "Evidencia que sirve","Estado","Responsable","Fecha de revisión","Observaciones y brecha detectada"]
heads(ws, HEAD, 5)
r0 = 6
prev = None
for i, (amb, que, art, apl, ev) in enumerate(R):
    r = r0 + i
    vals = [i+1, amb, que, art, apl, ev, "", "", "", ""]
    repetido = (amb == prev); prev = amb
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=2+j, value=v)
        c.border = box
        c.font = fo(9)
        c.alignment = Alignment(vertical="top", wrap_text=True,
                                horizontal="center" if j in (0,3,4) else "left")
        if j == 1:
            c.font = fo(9, not repetido, "0C1E50" if not repetido else "6B76A0")
        if j in (6,7,8,9):
            c.fill = PatternFill("solid", fgColor=BAND)
    ws.cell(row=r, column=10).number_format = "dd-mm-yyyy"
    ws.row_dimensions[r].height = 40
rN = r0 + len(R) - 1
dv = DataValidation(type="list", allow_blank=True, formula1='"Cumple,No cumple,No aplica,Pendiente"')
ws.add_data_validation(dv); dv.add("H%d:H%d" % (r0, rN))
rng = "H%d:H%d" % (r0, rN)
ws.conditional_formatting.add(rng, FormulaRule(formula=['$H%d="Cumple"' % r0], fill=PatternFill("solid", fgColor=OK), stopIfTrue=False))
ws.conditional_formatting.add(rng, FormulaRule(formula=['$H%d="No cumple"' % r0], fill=PatternFill("solid", fgColor=BAD), stopIfTrue=False))
ws.conditional_formatting.add(rng, FormulaRule(formula=['$H%d="No aplica"' % r0], fill=PatternFill("solid", fgColor=NA), stopIfTrue=False))
ws.freeze_panes = "D6"
ws.auto_filter.ref = "B5:K%d" % rN

AMB = []
for a,_,_,_,_ in R:
    if a not in AMB: AMB.append(a)

# ══════════════ HOJA 2 · RESUMEN ══════════════
ws2 = wb.create_sheet("Resumen")
ws2.sheet_view.showGridLines = False
for k, v in [("A",2.5),("B",26),("C",12),("D",12),("E",12),("F",12),("G",16),("H",44)]:
    ws2.column_dimensions[k].width = v
banda(ws2, "B2:H2", "RESUMEN DE CUMPLIMIENTO")
sub(ws2, "B3:H3", "Se calcula solo a partir de la hoja «Checklist». El porcentaje excluye las líneas marcadas «No aplica».")
heads(ws2, ["Ámbito","Cumple","No cumple","No aplica","Pendiente","% cumplimiento","Qué mirar si sale bajo"], 5, alto=34)
NOTA = {
 "Gestión general":"Suele fallar la difusión permanente del artículo 4 N°4: existe la matriz, pero no hay campañas registradas.",
 "Matriz (MIPER)":"Lo más observado es la metodología: la evaluación debe seguir la Guía Técnica del ISP, no un criterio propio.",
 "Programa preventivo":"Revisa primero alcohol y drogas, vida saludable y la evaluación anual de eficacia: son las que casi nunca están.",
 "Control de riesgos y EPP":"El EPP es el último recurso. Si no hay evidencia de medidas colectivas previas, el orden de prelación no se cumplió.",
 "Información y capacitación":"El IRL previo al inicio de labores y las 8 horas del curso son los dos datos duros que se piden.",
 "Participación":"El registro del acta en la Dirección del Trabajo dentro de 15 días hábiles es un incumplimiento con fecha.",
 "Riesgo grave y emergencias":"El simulacro anual es el que más se olvida y el más fácil de acreditar.",
 "Coordinación":"Una acreditación documental de contratistas no basta: hay que mostrar intercambio real de riesgos y medidas.",
 "Estructura preventiva":"Verifica la categoría y el tiempo de dedicación del experto, no solo que exista el Departamento.",
 "Reglamento interno":"Los 30 días de anticipación y el ingreso en ambas páginas web son los puntos que se caen.",
 "Mapas de riesgos":"Desde el 1 de marzo de 2025 las actualizaciones deben seguir la guía de la Resolución Exenta N°129.",
 "Vigilancia de la salud":"Si hay un agente de riesgo evaluado, tiene que existir la solicitud de incorporación al programa.",
 "Investigación":"Debe usarse la metodología del organismo administrador, con enfoque de género y participación.",
 "Registros e indicadores":"Casi todo se cae por la desagregación por sexo del artículo 74.",
}
r = 6
for a in AMB:
    ws2.cell(row=r, column=2, value=a).font = fo(9, True)
    for j, est in enumerate(["Cumple","No cumple","No aplica","Pendiente"]):
        ws2.cell(row=r, column=3+j,
                 value='=COUNTIFS(Checklist!$C$%d:$C$%d,$B%d,Checklist!$H$%d:$H$%d,"%s")'
                       % (r0, rN, r, r0, rN, est))
    ws2.cell(row=r, column=7, value='=IF((C%d+D%d)=0,"—",C%d/(C%d+D%d))' % (r,r,r,r,r))
    ws2.cell(row=r, column=7).number_format = "0%"
    ws2.cell(row=r, column=8, value=NOTA.get(a,"")).alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(2, 9):
        c = ws2.cell(row=r, column=col); c.border = box
        if c.font.size is None or col > 2: c.font = fo(9)
        c.alignment = Alignment(vertical="top", wrap_text=(col == 8),
                                horizontal="center" if 3 <= col <= 7 else "left")
        if 3 <= col <= 7: c.fill = PatternFill("solid", fgColor=CALC)
    ws2.cell(row=r, column=2).font = fo(9, True)
    ws2.row_dimensions[r].height = 30
    r += 1
tot = r
ws2.cell(row=tot, column=2, value="TOTAL").font = fo(10, True, "FFFFFF")
for j, est in enumerate(["Cumple","No cumple","No aplica","Pendiente"]):
    ws2.cell(row=tot, column=3+j, value='=SUM(%s6:%s%d)' % (chr(67+j), chr(67+j), tot-1))
ws2.cell(row=tot, column=7, value='=IF((C%d+D%d)=0,"—",C%d/(C%d+D%d))' % (tot,tot,tot,tot,tot))
ws2.cell(row=tot, column=7).number_format = "0%"
ws2.cell(row=tot, column=8, value="El porcentaje se calcula sobre las líneas efectivamente aplicables.")
for col in range(2, 9):
    c = ws2.cell(row=tot, column=col)
    c.fill = PatternFill("solid", fgColor=NAVY); c.border = box
    c.font = fo(10, True, "FFFFFF")
    c.alignment = Alignment(vertical="center", wrap_text=(col == 8),
                            horizontal="center" if 3 <= col <= 7 else "left")
ws2.row_dimensions[tot].height = 26
ws2.cell(row=tot+2, column=2,
         value="Sin líneas revisadas el porcentaje aparece como «—». No es un 0%: es que todavía no se ha evaluado nada.").font = fo(8.5, False, GREY)

# ══════════════ HOJA 3 · PLAN DE CIERRE ══════════════
ws3 = wb.create_sheet("Plan de cierre")
ws3.sheet_view.showGridLines = False
for k, v in [("A",2.5),("B",7),("C",12),("D",50),("E",50),("F",20),("G",14),("H",14),("I",30)]:
    ws3.column_dimensions[k].width = v
banda(ws3, "B2:I2", "PLAN DE CIERRE DE BRECHAS")
sub(ws3, "B3:I3", "Una fila por cada línea marcada «No cumple». Copia el N° del checklist para no perder la trazabilidad.")
heads(ws3, ["N°","N° checklist","Hallazgo","Acción correctiva comprometida","Responsable",
            "Plazo","Estado","Evidencia de cierre"], 5, alto=34)
for i in range(40):
    r = 6 + i
    ws3.cell(row=r, column=2, value=i+1)
    for col in range(2, 10):
        c = ws3.cell(row=r, column=col); c.border = box; c.font = fo(9)
        c.alignment = Alignment(vertical="top", wrap_text=True,
                                horizontal="center" if col in (2,3,7,8) else "left")
        if col > 3: c.fill = PatternFill("solid", fgColor=BAND)
    ws3.cell(row=r, column=7).number_format = "dd-mm-yyyy"
    ws3.row_dimensions[r].height = 30
dv3 = DataValidation(type="list", allow_blank=True, formula1='"Abierta,En curso,Cerrada"')
ws3.add_data_validation(dv3); dv3.add("H6:H45")
ws3.freeze_panes = "D6"

# ══════════════ HOJA 4 · CÓMO USARLA ══════════════
ws4 = wb.create_sheet("Cómo usarla")
ws4.sheet_view.showGridLines = False
for k, v in [("A",2.5),("B",26),("C",104)]:
    ws4.column_dimensions[k].width = v
banda(ws4, "B2:C2", "CÓMO USAR ESTE CHECKLIST")
FILAS = [
("Para qué sirve","Es una preauditoría: te dice, línea por línea, qué te pueden pedir en una fiscalización del DS 44 y qué evidencia sirve para responder. No reemplaza al Formulario Único de Fiscalización de la Dirección del Trabajo; te prepara para él."),
("De dónde sale cada línea","Del texto del Decreto Supremo N° 44 del Ministerio del Trabajo y Previsión Social, publicado en el Diario Oficial el 27 de julio de 2024 y vigente desde el 1 de febrero de 2025. La columna «Artículo DS 44» indica de dónde sale cada exigencia, para que puedas ir a la fuente."),
("Cómo se completa","1) Filtra por «A quién aplica» y marca «No aplica» en lo que no te corresponde por dotación o por actividad. 2) Recorre el resto marcando Cumple, No cumple o Pendiente. 3) Anota la evidencia concreta en «Observaciones». 4) Lleva cada «No cumple» a la hoja «Plan de cierre»."),
("Cómo se lee el resultado","La hoja «Resumen» se calcula sola. El porcentaje excluye las líneas «No aplica», porque no tiene sentido penalizar a una empresa por no tener Departamento de Prevención si no está obligada a tenerlo."),
("Dónde se cuenta cada umbral","Ojo con esto, porque cambia el resultado: el Comité Paritario y el Delegado se miden por empresa, faena, sucursal o agencia. El encargado capacitado y el Departamento de Prevención se miden por entidad empleadora. Usar una sola cifra corporativa para todo es el error más común."),
("Qué NO cubre","Las obligaciones que vienen de otras normas: protocolos del Ministerio de Salud, reglamentos sectoriales, la ley 16.744 y el Código del Trabajo. Este checklist cubre el DS 44."),
("Una advertencia honesta","Marcar «Cumple» sin evidencia no sirve de nada. El artículo 72 exige que toda la gestión esté registrada y respaldada de forma documental y fidedigna. Si no puedes mostrar el documento, para la fiscalización no ocurrió."),
("Quién lo hizo","Tazki, software chileno de gestión de seguridad y salud en el trabajo. Si quieres que estos registros dejen de vivir en carpetas y planillas, escríbenos en tazki.cl"),
]
r = 4
for t, d in FILAS:
    a = ws4.cell(row=r, column=2, value=t); a.font = fo(10, True); a.fill = PatternFill("solid", fgColor=SOFT)
    a.alignment = Alignment(vertical="top", wrap_text=True); a.border = box
    b = ws4.cell(row=r, column=3, value=d); b.font = fo(10)
    b.alignment = Alignment(vertical="top", wrap_text=True); b.border = box
    ws4.row_dimensions[r].height = 58
    r += 1

# ══════════════ HOJA 5 · CONTROL DE VERSIONES ══════════════
ws5 = wb.create_sheet("Control de versiones")
ws5.sheet_view.showGridLines = False
for k, v in [("A",2.5),("B",14),("C",16),("D",30),("E",56)]:
    ws5.column_dimensions[k].width = v
banda(ws5, "B2:E2", "CONTROL DE VERSIONES")
heads(ws5, ["Versión","Fecha","Responsable","Qué cambió"], 4, alto=26)
for i in range(15):
    r = 5 + i
    for col in range(2, 6):
        c = ws5.cell(row=r, column=col); c.border = box; c.font = fo(9)
        c.alignment = Alignment(vertical="top", wrap_text=True,
                                horizontal="center" if col in (2,3) else "left")
    ws5.cell(row=r, column=3).number_format = "dd-mm-yyyy"
    ws5.row_dimensions[r].height = 24

for s in wb.worksheets:
    s.sheet_properties.tabColor = COBALT
wb.save("/tmp/piezas/Checklist-DS44-Tazki.xlsx")
print("guardado. hojas:", [s.title for s in wb.worksheets], "· filas:", len(R), "· ámbitos:", len(AMB))
